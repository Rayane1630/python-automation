"""
OUTILS EXCEL ET FICHIERS
========================

Fonctions courtes et réutilisables pour automatiser les tâches Excel du
quotidien : nommer un fichier avec la date, créer un classeur propre,
ajouter un onglet, retrouver le dernier fichier reçu, consolider un dossier.

Exemple rapide
--------------
    from excel_outils import nom_fichier, creer_excel
    import pandas as pd

    df = pd.DataFrame({"Fonds": ["A", "B"], "NAV": [101.25, 98.40]})
    chemin = Path(r"XXXXXXXX") / nom_fichier("Controle_NAV")   # Controle_NAV_20261028.xlsx
    creer_excel(chemin, {"NAV": df})
"""

from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Style par défaut des en-têtes
COULEUR_ENTETE = "1F3864"
POLICE = "Arial"


# ---------------------------------------------------------------------------
# NOMS DE FICHIERS ET DOSSIERS
# ---------------------------------------------------------------------------

def nom_fichier(prefixe, jour=None, extension="xlsx", format_date="%Y%m%d"):
    """
    Construit un nom de fichier daté.
        nom_fichier("Controle_NAV")                      -> Controle_NAV_20261028.xlsx
        nom_fichier("Reporting", date(2026, 9, 30), "pdf") -> Reporting_20260930.pdf
    Le format AAAAMMJJ trie les fichiers dans l'ordre chronologique.
    """
    jour = jour or date.today()
    return f"{prefixe}_{jour.strftime(format_date)}.{extension.lstrip('.')}"


def dossier_du_jour(racine, jour=None):
    """Crée (si besoin) et renvoie un sous-dossier daté : racine/2026/10/20261028."""
    jour = jour or date.today()
    dossier = Path(racine) / f"{jour:%Y}" / f"{jour:%m}" / f"{jour:%Y%m%d}"
    dossier.mkdir(parents=True, exist_ok=True)
    return dossier


def dernier_fichier(dossier, motif="*.xlsx"):
    """Renvoie le fichier le plus récent d'un dossier (par date de modification)."""
    fichiers = sorted(Path(dossier).glob(motif), key=lambda f: f.stat().st_mtime)
    if not fichiers:
        raise FileNotFoundError(f"Aucun fichier '{motif}' dans {dossier}")
    return fichiers[-1]


def renommer_avec_date(chemin, jour=None):
    """Ajoute la date au nom d'un fichier existant : Export.xlsx -> Export_20261028.xlsx."""
    chemin = Path(chemin)
    nouveau = chemin.with_name(nom_fichier(chemin.stem, jour, chemin.suffix))
    chemin.rename(nouveau)
    return nouveau


# ---------------------------------------------------------------------------
# LECTURE
# ---------------------------------------------------------------------------

def lire_excel(chemin, onglet=0, ligne_entete=1):
    """
    Lit un onglet dans un DataFrame.
    ligne_entete : numéro de ligne Excel des en-têtes (1 = première ligne).
    """
    return pd.read_excel(chemin, sheet_name=onglet, header=ligne_entete - 1)


def consolider_dossier(dossier, motif="*.xlsx", onglet=0):
    """
    Empile tous les fichiers d'un dossier dans un seul tableau,
    avec une colonne "Fichier_source" pour savoir d'où vient chaque ligne.
    Accepte aussi les .csv (séparateur ; ou , détecté automatiquement).
    """
    tables = []
    for f in sorted(Path(dossier).glob(motif)):
        if f.name.startswith("~$"):              # fichiers temporaires d'Excel ouverts
            continue
        if f.suffix.lower() == ".csv":
            df = pd.read_csv(f, sep=None, engine="python")
        else:
            df = pd.read_excel(f, sheet_name=onglet)
        df.insert(0, "Fichier_source", f.name)
        tables.append(df)
    if not tables:
        raise FileNotFoundError(f"Aucun fichier '{motif}' dans {dossier}")
    return pd.concat(tables, ignore_index=True)


# ---------------------------------------------------------------------------
# ÉCRITURE ET MISE EN FORME
# ---------------------------------------------------------------------------

def mettre_en_forme(chemin, formats_colonnes=None):
    """
    Met en forme tous les onglets d'un classeur :
        - en-têtes en gras, fond bleu foncé, texte blanc ;
        - volets figés sous la ligne d'en-tête ;
        - filtre automatique ;
        - largeur des colonnes ajustée au contenu ;
        - formats de nombres optionnels : {"NAV": "#,##0.00", "Poids": "0.00%"}.
    """
    formats_colonnes = formats_colonnes or {}
    classeur = load_workbook(chemin)
    trait = Side(style="thin", color="BFBFBF")

    for feuille in classeur.worksheets:
        if feuille.max_row == 0:
            continue
        for cellule in feuille[1]:
            cellule.font = Font(name=POLICE, bold=True, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor=COULEUR_ENTETE)
            cellule.alignment = Alignment(horizontal="center", vertical="center")

        entetes = {c.value: c.column for c in feuille[1]}
        for ligne in feuille.iter_rows(min_row=2):
            for cellule in ligne:
                cellule.font = Font(name=POLICE)
                cellule.border = Border(bottom=trait)
                nom_col = feuille.cell(row=1, column=cellule.column).value
                if nom_col in formats_colonnes:
                    cellule.number_format = formats_colonnes[nom_col]
                elif isinstance(cellule.value, (datetime, date)):
                    cellule.number_format = "dd/mm/yyyy"

        for nom_col, num_col in entetes.items():
            lettre = get_column_letter(num_col)
            longueur = max(len(str(c.value)) if c.value is not None else 0
                           for c in feuille[lettre])
            feuille.column_dimensions[lettre].width = min(max(10, longueur + 2), 60)

        feuille.freeze_panes = "A2"
        feuille.auto_filter.ref = feuille.dimensions

    classeur.save(chemin)


def creer_excel(chemin, onglets, formats_colonnes=None):
    """
    Crée un classeur avec un onglet par DataFrame, déjà mis en forme.
        onglets = {"NAV": df_nav, "Ratios": df_ratios}
    Crée le dossier parent si besoin. Renvoie le chemin du fichier.
    """
    chemin = Path(chemin)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(chemin, engine="openpyxl") as writer:
        for nom, df in onglets.items():
            df.to_excel(writer, sheet_name=nom[:31], index=False)   # 31 caractères max
    mettre_en_forme(chemin, formats_colonnes)
    return chemin


def ajouter_onglet(chemin, nom, df, formats_colonnes=None):
    """Ajoute (ou remplace) un onglet dans un classeur existant."""
    with pd.ExcelWriter(chemin, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        df.to_excel(writer, sheet_name=nom[:31], index=False)
    mettre_en_forme(chemin, formats_colonnes)
    return Path(chemin)
